# Copyright (c) 2026 芥子星辰（杭州）智能科技有限公司
# Licensed under the SoloBooks License v2.0.

"""备份与归档（P0 WP5）：一键导出全账套归档 Excel。

定位：用户问"我的数据能不能拿走"时，交给他一份能直接给会计/审计的工作簿。
封面带数据指纹（期间 + 凭证数 + 借贷合计），便于核对"导出的就是系统里的"。

有意裁剪（见 P0 计划）：不做加密备份包、不做定时邮件备份、不做"上传恢复覆盖账套"。
归档 Excel + 平台数据库自身备份已闭环，破坏性恢复能力超出当前需要。
"""
from __future__ import annotations

import calendar
import hashlib
from datetime import date, datetime, timezone
from decimal import Decimal
from pathlib import Path
from typing import Any, Iterable
from uuid import uuid4

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from sqlalchemy import select

from core.money import ZERO, money, money_text
from db.models import BackupRecord, ChartAccount, Invoice, JournalEntry, JournalLine
from db.session import session_scope
from services.account_service import get_company

ARCHIVE_DIR = Path(__file__).resolve().parent.parent / "data" / "archives"
MAX_MONTHS = 24
MAX_LINES = 50_000

_HEADER_FONT = Font(bold=True, size=11)
_HEADER_FILL = PatternFill("solid", fgColor="EAF1FB")
_TITLE_FONT = Font(bold=True, size=16)
_THIN = Side(style="thin", color="B9C4D2")
_BORDER = Border(left=_THIN, right=_THIN, top=_THIN, bottom=_THIN)


def _month_list(year: int, months: str | None) -> list[str]:
    if months:
        text = str(months).strip()
        try:
            if "-" in text:
                start, end = (int(part) for part in text.split("-", 1))
            else:
                start = end = int(text)
        except ValueError as exc:
            raise ValueError("期间格式应为 1-10 或单个月份 3") from exc
        if not (1 <= start <= 12 and 1 <= end <= 12 and start <= end):
            raise ValueError("期间范围应在 1-12 之间且起始不晚于结束")
        return [f"{year}-{month:02d}" for month in range(start, end + 1)]
    return [f"{year}-{month:02d}" for month in range(1, 13)]


def _month_range(month: str) -> tuple[date, date]:
    year, number = (int(part) for part in month.split("-", 1))
    return date(year, number, 1), date(year, number, calendar.monthrange(year, number)[1])


def _sheet(wb: Workbook, title: str, headers: list[str], widths: list[int] | None = None):
    ws = wb.create_sheet(title)
    ws.append(headers)
    for index, cell in enumerate(ws[1], start=1):
        cell.font = _HEADER_FONT
        cell.fill = _HEADER_FILL
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = _BORDER
    ws.freeze_panes = "A2"
    if widths:
        for index, width in enumerate(widths, start=1):
            ws.column_dimensions[get_column_letter(index)].width = width
    return ws


def _style_rows(ws, *, money_columns: Iterable[int] = ()) -> None:
    money_set = set(money_columns)
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.border = _BORDER
            if cell.column in money_set and isinstance(cell.value, (int, float, Decimal)):
                cell.number_format = "#,##0.00"
                cell.alignment = Alignment(horizontal="right")


async def _load_entries(company_id: str, start: date, end: date) -> list[tuple[JournalEntry, list[JournalLine]]]:
    async with session_scope() as session:
        entries = (
            await session.execute(
                select(JournalEntry)
                .where(
                    JournalEntry.company_id == company_id,
                    JournalEntry.status.in_(("posted", "reversed")),
                    JournalEntry.period_start >= start,
                    JournalEntry.period_start <= end,
                )
                .order_by(JournalEntry.period_start, JournalEntry.voucher_no, JournalEntry.occurred_on)
            )
        ).scalars().all()
        ids = [entry.id for entry in entries]
        lines = (
            await session.execute(
                select(JournalLine)
                .where(JournalLine.journal_entry_id.in_(ids))
                .order_by(JournalLine.journal_entry_id, JournalLine.line_no)
            )
        ).scalars().all() if ids else []
    grouped: dict[str, list[JournalLine]] = {}
    for line in lines:
        grouped.setdefault(line.journal_entry_id, []).append(line)
    return [(entry, grouped.get(entry.id, [])) for entry in entries]


async def _opening_by_account(company_id: str, start: date) -> dict[str, Decimal]:
    """期间开始前的累计净额（借余为正），明细账的期初行用它。"""
    async with session_scope() as session:
        rows = (
            await session.execute(
                select(JournalLine.account_code, JournalLine.direction, JournalLine.amount)
                .join(JournalEntry, JournalEntry.id == JournalLine.journal_entry_id)
                .where(
                    JournalEntry.company_id == company_id,
                    JournalEntry.status.in_(("posted", "reversed")),
                    JournalEntry.period_start < start,
                )
            )
        ).all()
    balances: dict[str, Decimal] = {}
    for code, direction, amount in rows:
        signed = money(amount) if direction == "debit" else -money(amount)
        balances[code] = balances.get(code, ZERO) + signed
    return balances


async def _load_invoices(company_id: str, start: date, end: date) -> list[Invoice]:
    async with session_scope() as session:
        return list(
            (
                await session.execute(
                    select(Invoice)
                    .where(
                        Invoice.company_id == company_id,
                        Invoice.invoice_date >= start,
                        Invoice.invoice_date <= end,
                    )
                    .order_by(Invoice.invoice_date, Invoice.invoice_number)
                )
            ).scalars().all()
        )


async def _load_salary(company_id: str, start: date, end: date) -> list[dict[str, Any]]:
    """工资表：legacy sqlite 表，表不存在时静默返回空（归档不能因可选模块失败）。"""
    from services.database import get_conn

    try:
        conn = await get_conn()
        try:
            cursor = await conn.execute(
                "SELECT r.employee_name, s.month, r.base_salary, r.bonus, r.allowance, r.deduction,"
                " r.personal_social_security, r.personal_tax, r.company_social_security,"
                " r.net_salary, r.total_cost"
                " FROM salary_records r JOIN salary_sheets s ON s.id = r.sheet_id"
                " WHERE s.company_id = ? AND s.month >= ? AND s.month <= ?"
                " ORDER BY s.month, r.employee_name",
                (company_id, start.strftime("%Y-%m"), end.strftime("%Y-%m")),
            )
            rows = await cursor.fetchall()
        finally:
            await conn.close()
    except Exception:
        return []
    return [
        {
            "employee": row[0],
            "month": row[1],
            "base": row[2],
            "bonus": row[3],
            "allowance": row[4],
            "deduction": row[5],
            "social": row[6],
            "tax": row[7],
            "companySocial": row[8],
            "net": row[9],
            "cost": row[10],
        }
        for row in rows
    ]


async def build_archive(
    company_id: str,
    user_id: str,
    *,
    year: int,
    months: str | None = None,
) -> dict[str, Any]:
    month_list = _month_list(year, months)
    if len(month_list) > MAX_MONTHS:
        raise ValueError(f"一次最多归档 {MAX_MONTHS} 个月，请分批导出")
    start, _ = _month_range(month_list[0])
    _, end = _month_range(month_list[-1])

    company = await get_company(company_id) or {}
    company_name = str(company.get("name") or company.get("company_name") or "账套")
    entries = await _load_entries(company_id, start, end)
    line_count = sum(len(lines) for _, lines in entries)
    if line_count > MAX_LINES:
        raise ValueError(
            f"本次范围包含 {line_count} 条分录行，超过 {MAX_LINES} 条上限，请按月分批归档"
        )

    debit_total = sum(
        (money(line.amount) for _, lines in entries for line in lines if line.direction == "debit"), ZERO
    )
    credit_total = sum(
        (money(line.amount) for _, lines in entries for line in lines if line.direction == "credit"), ZERO
    )
    fingerprint = hashlib.sha256(
        f"{month_list[0]}:{month_list[-1]}:{len(entries)}:{line_count}:{debit_total}:{credit_total}".encode()
    ).hexdigest()[:32]

    wb = Workbook()
    wb.remove(wb.active)
    _write_cover(wb, company_name, company, month_list, len(entries), line_count, debit_total, credit_total, fingerprint, user_id)
    _write_account_balances(wb, await _load_balance_summary(company_id, month_list[-1]))
    _write_voucher_list(wb, entries)
    await _write_sub_ledger(wb, company_id, start, entries)
    _write_invoices(wb, await _load_invoices(company_id, start, end))
    _write_fixed_assets(wb, await _load_fixed_assets(user_id, company_id))
    _write_salary(wb, await _load_salary(company_id, start, end))
    await _write_statements(wb, user_id, company_id, month_list[-1], company_name)

    ARCHIVE_DIR.mkdir(parents=True, exist_ok=True)
    safe_name = "".join(ch for ch in company_name if ch not in '\\/:*?"<>|').strip() or "账套"
    file_name = f"{safe_name}_归档_{month_list[0]}_{month_list[-1]}_{datetime.now():%Y%m%d%H%M%S}.xlsx"
    path = ARCHIVE_DIR / f"{company_id[:12]}_{uuid4().hex[:8]}_{file_name}"
    wb.save(path)
    size = path.stat().st_size

    async with session_scope() as session:
        record = BackupRecord(
            id=str(uuid4()),
            company_id=company_id,
            file_path=str(path),
            file_name=file_name,
            scope=f"{month_list[0]}~{month_list[-1]}",
            period_from=month_list[0],
            period_to=month_list[-1],
            size_bytes=size,
            voucher_count=len(entries),
            data_fingerprint=fingerprint,
            created_by=user_id,
        )
        session.add(record)
        await session.flush()
        record_id = record.id

    return {
        "id": record_id,
        "fileName": file_name,
        "filePath": str(path),
        "sizeBytes": size,
        "voucherCount": len(entries),
        "lineCount": line_count,
        "periodFrom": month_list[0],
        "periodTo": month_list[-1],
        "fingerprint": fingerprint,
        "debitTotal": money_text(debit_total),
        "creditTotal": money_text(credit_total),
    }


def _write_cover(wb, company_name, company, month_list, voucher_count, line_count, debit_total, credit_total, fingerprint, user_id):
    ws = wb.create_sheet("封面")
    ws.column_dimensions["A"].width = 18
    ws.column_dimensions["B"].width = 52
    ws["A1"] = "会 计 档 案 归 档"
    ws["A1"].font = _TITLE_FONT
    ws.merge_cells("A1:B1")
    ws["A1"].alignment = Alignment(horizontal="center", vertical="center")
    rows = [
        ("账套名称", company_name),
        ("纳税人识别号", company.get("taxId") or company.get("tax_id") or ""),
        ("归档期间", f"{month_list[0]} 至 {month_list[-1]}"),
        ("凭证张数", voucher_count),
        ("分录行数", line_count),
        ("借方合计", money_text(debit_total)),
        ("贷方合计", money_text(credit_total)),
        ("借贷平衡", "是" if debit_total == credit_total else "否（请先核对账务）"),
        ("数据指纹", fingerprint),
        ("生成时间", datetime.now().strftime("%Y-%m-%d %H:%M:%S")),
        ("生成人", user_id),
        ("说明", "本工作簿由梭罗记账自动生成，封面数据指纹用于核对导出内容与系统账务一致。"),
    ]
    for index, (label, value) in enumerate(rows, start=3):
        ws.cell(row=index, column=1, value=label).font = Font(bold=True)
        ws.cell(row=index, column=2, value=value)
        ws.cell(row=index, column=1).border = _BORDER
        ws.cell(row=index, column=2).border = _BORDER


async def _load_balance_summary(company_id: str, month: str) -> dict[str, Any]:
    """科目余额表直接复用整册口径，保证归档与页面同数。"""
    from services.ledger_book_service import build_ledger_book

    try:
        book = await build_ledger_book(company_id, month=month, include_zero_rows=True)
        return book["balance"]
    except Exception:
        return {"items": []}


async def _load_fixed_assets(user_id: str, company_id: str) -> list[dict[str, Any]]:
    """固定资产（P1 WP-B 起为核心 ORM）；读失败不影响归档其余部分。"""
    try:
        from sqlalchemy import select

        from core.fixed_asset_categories import get_category
        from core.money import ZERO
        from db.models import FixedAsset
        from db.session import session_scope

        async with session_scope() as session:
            rows = (
                await session.execute(
                    select(FixedAsset).where(
                        FixedAsset.company_id == company_id,
                        FixedAsset.deleted_at.is_(None),
                    )
                )
            ).scalars().all()
            return [
                {
                    "name": row.name,
                    "category": get_category(row.category).name,
                    "originalValue": row.original_value,
                    "startDate": row.available_for_use_date.isoformat() if row.available_for_use_date else "",
                    "usefulLife": row.useful_life_months,
                    "residualRate": row.residual_rate,
                    "accumulatedDepreciation": row.accumulated_depreciation or ZERO,
                    "netValue": row.original_value - (row.accumulated_depreciation or ZERO),
                    "status": row.status,
                }
                for row in rows
            ]
    except Exception:
        return []


def _write_account_balances(wb, summary: dict[str, Any]) -> None:
    ws = _sheet(
        wb,
        "科目余额表",
        ["科目编码", "科目名称", "期初余额", "借方发生", "贷方发生", "期末余额"],
        [14, 34, 16, 16, 16, 16],
    )
    for item in summary.get("items", []):
        ws.append(
            [
                item["accountCode"],
                item["accountName"] + ("（子级合计）" if item.get("isParent") else ""),
                float(money(item["openingBalance"])),
                float(money(item["debitTotal"])),
                float(money(item["creditTotal"])),
                float(money(item["closingBalance"])),
            ]
        )
    _style_rows(ws, money_columns=(3, 4, 5, 6))


def _write_voucher_list(wb, entries) -> None:
    ws = _sheet(
        wb,
        "凭证清单",
        ["期间", "凭证号", "日期", "摘要", "科目编码", "科目名称", "方向", "金额", "往来单位", "部门", "项目"],
        [10, 20, 12, 34, 14, 30, 8, 16, 24, 16, 16],
    )
    for entry, lines in entries:
        for line in lines:
            auxiliary = line.auxiliary if isinstance(line.auxiliary, dict) else {}
            ws.append(
                [
                    entry.period_start.strftime("%Y-%m"),
                    entry.voucher_no or "",
                    entry.occurred_on.isoformat(),
                    entry.summary,
                    line.account_code,
                    line.account_name,
                    "借" if line.direction == "debit" else "贷",
                    float(money(line.amount)),
                    auxiliary.get("partner_name") or "",
                    auxiliary.get("department_name") or "",
                    auxiliary.get("project_name") or "",
                ]
            )
    _style_rows(ws, money_columns=(8,))


async def _write_sub_ledger(wb, company_id: str, start: date, entries) -> None:
    opening = await _opening_by_account(company_id, start)
    ws = _sheet(
        wb,
        "明细账",
        ["科目编码", "科目名称", "日期", "凭证号", "摘要", "借方", "贷方", "余额"],
        [14, 30, 12, 20, 36, 16, 16, 16],
    )
    running: dict[str, Decimal] = dict(opening)
    accounts: dict[str, str] = {}
    for entry, lines in entries:
        for line in lines:
            accounts.setdefault(line.account_code, line.account_name)
    # 期初行
    for code in sorted(accounts):
        balance = running.get(code, ZERO)
        if balance:
            ws.append([code, accounts[code], "", "", "期初余额", "", "", float(balance)])
    for entry, lines in entries:
        for line in lines:
            signed = money(line.amount) if line.direction == "debit" else -money(line.amount)
            running[line.account_code] = running.get(line.account_code, ZERO) + signed
            ws.append(
                [
                    line.account_code,
                    line.account_name,
                    entry.occurred_on.isoformat(),
                    entry.voucher_no or "",
                    entry.summary,
                    float(money(line.amount)) if line.direction == "debit" else "",
                    float(money(line.amount)) if line.direction == "credit" else "",
                    float(running[line.account_code]),
                ]
            )
    _style_rows(ws, money_columns=(6, 7, 8))


def _write_invoices(wb, invoices: list[Invoice]) -> None:
    ws = _sheet(
        wb,
        "进销项发票",
        ["类型", "开票日期", "发票号码", "销方", "购方", "不含税金额", "税额", "价税合计", "项目", "备注"],
        [10, 12, 22, 30, 30, 16, 14, 16, 26, 30],
    )
    for invoice in invoices:
        ws.append(
            [
                "销项" if invoice.direction == "income" else "进项",
                invoice.invoice_date.isoformat() if invoice.invoice_date else "",
                invoice.invoice_number,
                invoice.seller_name,
                invoice.buyer_name,
                float(money(invoice.amount_excl_tax)),
                float(money(invoice.tax_amount)),
                float(money(invoice.amount_incl_tax)),
                invoice.item_name,
                invoice.remark,
            ]
        )
    _style_rows(ws, money_columns=(6, 7, 8))


def _write_fixed_assets(wb, assets: list[dict[str, Any]]) -> None:
    ws = _sheet(
        wb,
        "固定资产",
        ["资产名称", "类别", "原值", "启用日期", "使用年限(月)", "残值率", "累计折旧", "净值", "状态"],
        [26, 16, 16, 12, 12, 10, 14, 14, 12],
    )
    status_labels = {
        "draft": "草稿", "pending_review": "待复核", "active": "使用中", "idle": "闲置",
        "fully_depreciated": "已折旧完毕", "disposing": "处置中", "disposed": "已处置", "scrapped": "已报废",
    }
    for asset in assets:
        ws.append(
            [
                asset.get("name") or "",
                asset.get("category") or "",
                float(money(asset.get("originalValue") or 0)),
                str(asset.get("startDate") or "")[:10],
                asset.get("usefulLife") or "",
                float(asset.get("residualRate") or 0),
                float(money(asset.get("accumulatedDepreciation") or 0)),
                float(money(asset.get("netValue") or 0)),
                status_labels.get(asset.get("status") or "", asset.get("status") or ""),
            ]
        )
    _style_rows(ws, money_columns=(3, 7, 8))


def _write_salary(wb, rows: list[dict[str, Any]]) -> None:
    ws = _sheet(
        wb,
        "工资表",
        ["月份", "员工", "基本工资", "奖金", "补贴", "扣款", "个人社保", "个税", "单位社保", "实发", "总成本"],
        [10, 16, 14, 12, 12, 12, 14, 12, 14, 14, 14],
    )
    for row in rows:
        ws.append(
            [
                row["month"],
                row["employee"],
                float(money(row["base"] or 0)),
                float(money(row["bonus"] or 0)),
                float(money(row["allowance"] or 0)),
                float(money(row["deduction"] or 0)),
                float(money(row["social"] or 0)),
                float(money(row["tax"] or 0)),
                float(money(row["companySocial"] or 0)),
                float(money(row["net"] or 0)),
                float(money(row["cost"] or 0)),
            ]
        )
    _style_rows(ws, money_columns=(3, 4, 5, 6, 7, 8, 9, 10, 11))


async def _write_statements(wb, user_id: str, company_id: str, month: str, company_name: str) -> None:
    ws = _sheet(wb, "三大报表", ["项目", "本期金额"], [40, 18])
    ws["A1"] = f"{company_name} · {month}"
    try:
        from services.tax_filing_export_service import build_statement_data

        statement = await build_statement_data(user_id, company_id, month, "month")
    except Exception:
        ws.append(["报表数据不可用（请检查账务是否已过账）", ""])
        return
    balance = statement.balance_current or {}
    profit = statement.profit_current or {}
    cash = statement.cash_current or {}

    def block(title: str, rows: list[tuple[str, Any]]) -> None:
        ws.append([])
        cell = ws.cell(row=ws.max_row + 1, column=1, value=title)
        cell.font = Font(bold=True, size=12)
        for label, value in rows:
            ws.append([label, float(money(value or 0))])

    block(
        "资产负债表",
        [
            ("货币资金", balance.get("cash")),
            ("应收账款", balance.get("accounts_receivable")),
            ("其他应收款", balance.get("other_receivable")),
            ("存货", balance.get("inventory")),
            ("固定资产原价", balance.get("fixed_asset_cost")),
            ("累计折旧", balance.get("accumulated_depreciation")),
            ("资产总计", balance.get("asset_total")),
            ("应付账款", balance.get("accounts_payable")),
            ("应付职工薪酬", balance.get("salary_payable")),
            ("应交税费", balance.get("tax_payable")),
            ("负债合计", balance.get("liability_total")),
            ("所有者权益合计", balance.get("equity_total")),
        ],
    )
    block(
        "利润表（本期）",
        [
            ("营业收入", profit.get("revenue")),
            ("营业成本", profit.get("cost")),
            ("管理费用", profit.get("management_expense")),
            ("销售费用", profit.get("sales_expense")),
            ("财务费用", profit.get("finance_expense")),
            ("利润总额", profit.get("total_profit")),
            ("净利润", profit.get("net_profit")),
        ],
    )
    block(
        "现金流量表（本期）",
        [
            ("销售商品收到的现金", cash.get("sales_cash")),
            ("经营活动现金流入", cash.get("operating_inflow")),
            ("经营活动现金流出", cash.get("operating_outflow")),
            ("现金净增加额", cash.get("net_increase")),
        ],
    )
    _style_rows(ws, money_columns=(2,))


async def list_records(company_id: str, limit: int = 50) -> list[dict[str, Any]]:
    async with session_scope() as session:
        rows = (
            await session.execute(
                select(BackupRecord)
                .where(BackupRecord.company_id == company_id)
                .order_by(BackupRecord.created_at.desc())
                .limit(limit)
            )
        ).scalars().all()
    return [
        {
            "id": row.id,
            "fileName": row.file_name,
            "scope": row.scope,
            "periodFrom": row.period_from,
            "periodTo": row.period_to,
            "sizeBytes": row.size_bytes,
            "voucherCount": row.voucher_count,
            "fingerprint": row.data_fingerprint,
            "createdBy": row.created_by,
            "createdAt": row.created_at.isoformat() if row.created_at else "",
            "exists": Path(row.file_path).exists(),
        }
        for row in rows
    ]


async def get_record(company_id: str, record_id: str) -> dict[str, Any] | None:
    async with session_scope() as session:
        row = await session.get(BackupRecord, record_id)
    if row is None or row.company_id != company_id:
        return None
    return {"id": row.id, "fileName": row.file_name, "filePath": row.file_path}
