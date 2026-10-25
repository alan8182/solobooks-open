# Copyright (c) 2026 芥子星辰（杭州）智能科技有限公司
# Licensed under the SoloBooks License v2.0.

"""备份归档（P0 WP5）：工作簿结构、数据指纹、留痕记录。

锁定三条口径：
1. 归档簿包含封面/科目余额表/凭证清单/明细账/进销项发票/固定资产/工资表/三大报表；
2. 封面指纹随数据变化（期间+凭证数+借贷合计）；
3. 每次导出留痕，可重新下载。
"""
from __future__ import annotations

from pathlib import Path

import pytest
from openpyxl import load_workbook

from services.accounting_core_service import create_entry, ensure_tenant
from services.backup_archive_service import build_archive, get_record, list_records

EXPECTED_SHEETS = {
    "封面",
    "科目余额表",
    "凭证清单",
    "明细账",
    "进销项发票",
    "固定资产",
    "工资表",
    "三大报表",
}


async def _ensure_tables() -> None:
    from db.base import Base
    import db.models  # noqa: F401
    from db.session import get_engine

    engine = get_engine()
    async with engine.begin() as connection:
        await connection.run_sync(Base.metadata.create_all)


async def _setup_company(company_id: str) -> None:
    from db.session import session_scope

    await _ensure_tables()
    async with session_scope() as session:
        await ensure_tenant(session, company_id, "owner-1")


async def _post_entry(company_id: str, entry_id: str, date: str, amount: str) -> None:
    from db.session import session_scope

    async with session_scope() as session:
        await create_entry(
            session,
            {
                "id": entry_id,
                "direction": "expense",
                "amount": amount,
                "title": "办公费",
                "category": "办公费",
                "occurDate": date,
                "sourceType": "manual",
                "voucherLines": {
                    "debit": [{"accountCode": "5602004", "accountName": "管理费用-办公费", "amount": amount}],
                    "credit": [{"accountCode": "1002", "accountName": "银行存款", "amount": amount}],
                },
            },
            company_id=company_id,
            user_id="owner-1",
            request_id=f"{entry_id}-req",
            confirm=True,
        )


@pytest.mark.asyncio
async def test_archive_workbook_structure_and_records():
    company_id = "c-backup-1"
    await _setup_company(company_id)
    await _post_entry(company_id, "backup-e1", "2026-10-08", "120.00")
    await _post_entry(company_id, "backup-e2", "2026-10-09", "80.00")

    result = await build_archive(company_id, "owner-1", year=2026, months="10-10")
    path = Path(result["filePath"])
    assert path.exists()
    assert result["voucherCount"] == 2
    assert result["periodFrom"] == "2026-10" and result["periodTo"] == "2026-10"
    assert result["fingerprint"]

    workbook = load_workbook(path)
    assert EXPECTED_SHEETS <= set(workbook.sheetnames)

    cover = workbook["封面"]
    values = {str(row[0].value): row[1].value for row in cover.iter_rows(min_row=3, max_col=2) if row[0].value}
    assert values["凭证张数"] == 2
    assert values["借贷平衡"] == "是"
    assert values["数据指纹"] == result["fingerprint"]

    vouchers = workbook["凭证清单"]
    # 表头 1 行 + 2 张凭证 × 2 条分录
    assert vouchers.max_row == 5

    records = await list_records(company_id)
    assert records and records[0]["id"] == result["id"]
    assert records[0]["exists"] is True

    record = await get_record(company_id, result["id"])
    assert record and record["filePath"] == str(path)


@pytest.mark.asyncio
async def test_archive_fingerprint_changes_with_data():
    company_id = "c-backup-fingerprint"
    await _setup_company(company_id)
    await _post_entry(company_id, "backup-f1", "2026-10-08", "100.00")
    first = await build_archive(company_id, "owner-1", year=2026, months="10-10")
    await _post_entry(company_id, "backup-f2", "2026-10-09", "100.00")
    second = await build_archive(company_id, "owner-1", year=2026, months="10-10")
    assert first["fingerprint"] != second["fingerprint"]
    assert second["voucherCount"] == 2


@pytest.mark.asyncio
async def test_archive_rejects_invalid_period():
    company_id = "c-backup-invalid"
    await _setup_company(company_id)
    with pytest.raises(ValueError):
        await build_archive(company_id, "owner-1", year=2026, months="13-14")
